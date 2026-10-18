"""Add the coach's expanded study list to data/manifest.json without touching the original topics.

usage: python3 tools/extend_manifest.py <Study_List.xlsx> [--write]

- Reads every tab of the spreadsheet. Each column is recognized by its CONTENT (not its header,
  since headers can be mislabeled), grouped into master lists, and run through tools/curation.py.
- Keeps all existing manifest topics exactly as they are (id, name, category, qbCategory, source);
  only their display `order` is renumbered so new topics can follow them inside each area.
- Appends new topics (with `aka` = other spellings from the spreadsheet, for search) and new areas.
- Rebuilds the anonymous team study lists from the tabs (identical tabs collapse into one list).
- Writes research/new_topics_queue.json: new topic ids in build-priority order (most-asked first,
  round-robin across areas so every area fills in early).
No student names are read into any output: tabs are only ever referred to by their contents.
"""
import json, re, sys, unicodedata
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import curation as C  # noqa: E402

MANIFEST = ROOT / "data" / "manifest.json"
BASE = ROOT / "research" / "manifest_base.json"   # the original 215-topic manifest (never changes)
QUEUE = ROOT / "research" / "new_topics_queue.json"
REPORT = ROOT / "research" / "curation_report.md"

AREAS = OrderedDict([
    ("American History", {"qb": "History · American History", "short": "US History"}),
    ("European History", {"qb": "History · European History", "short": "Euro History"}),
    ("Literature", {"qb": "Literature", "short": "Literature"}),
    ("Mythology", {"qb": "Mythology", "short": "Mythology"}),
    ("Religion", {"qb": "Religion", "short": "Religion"}),
    ("Philosophy", {"qb": "Philosophy", "short": "Philosophy"}),
    ("Social Science", {"qb": "Social Science", "short": "Social Science"}),
    ("Fine Arts", {"qb": "Fine Arts", "short": "Fine Arts"}),
    ("Music", {"qb": "Fine Arts · Auditory Fine Arts", "short": "Music"}),
    ("Physics", {"qb": "Science · Physics", "short": "Physics"}),
    ("Chemistry", {"qb": "Science · Chemistry", "short": "Chemistry"}),
    ("Biology", {"qb": "Science · Biology", "short": "Biology"}),
    ("Earth Science", {"qb": "Science · Other Science (Earth Science)", "short": "Earth Science"}),
    ("Astronomy", {"qb": "Science · Other Science (Astronomy)", "short": "Astronomy"}),
    ("Math", {"qb": "Science · Other Science (Math)", "short": "Math"}),
    ("Computer Science", {"qb": "Science · Other Science (Computer Science)", "short": "Comp Sci"}),
])

LIST_LABEL = {
    "us": "American History", "eu": "European History", "lit": "Literature",
    "litc": "Literature (mostly classical)", "myth": "Mythology", "rel": "Religion", "phil": "Philosophy",
    "soc": "Social Science", "fa": "Fine Arts", "mus": "Music", "sci": "All Science",
    "astro": "Science (mostly astronomy)", "biomath": "Biology + Math", "earth": "Earth Science",
}


def norm(s):
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[’'`.]", "", s)
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return re.sub(r"^the ", "", s)


def slug(s):
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[’'`.]", "", s)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def display(raw):
    raw = raw.strip().replace("'", "’")
    raw = re.sub(r'"([^"]*)"', r"“\1”", raw)
    if raw and raw[0].islower():
        raw = raw[0].upper() + raw[1:]
    return raw


AKA_JUNK = re.compile(r"\((?:\d{4} film|Decca album|Kiss video)\)|[a-z][A-Z]|^(?:Ronald|Herbert|Kennedy|Ivan|Diary|Alice|Night|Romeo|Awakening|Republic)$")


def aka_ok(name, alt):
    return alt != name and not AKA_JUNK.search(alt)


def parse_rules(text):
    rules = {}
    for line in text.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        if "=>" in line:
            raw, dec = (x.strip() for x in line.split("=>", 1))
        else:
            raw, dec = line, None
        rules[norm(raw)] = (raw, dec)
    return rules


def column_key(items):
    first = items[0]
    if first == "Hamlet":
        return "litc" if len(items) > 20 and items[20] == "Aristophanes" else "lit"
    if first == "Gravity":
        head = items[:25]
        return "astro" if "Jupiter" in head and "Prime number" not in head else "sci"
    by_first = {"George Washington": "us", "Martin Luther": "eu", "Odin": "myth", "Sikhism": "rel",
                "Plato": "phil", "Sigmund Freud": "soc", "Parthenon": "fa", "Johannes Brahms": "mus",
                "mitochondria": "biomath", "glaciers": "earth"}
    if first in by_first:
        return by_first[first]
    sys.exit(f"Unrecognized column starting with {first!r}")


def read_tabs(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tabs = []
    for ws in wb.worksheets:
        cols = []
        for c in range(1, ws.max_column + 1):
            items = [ws.cell(r, c).value for r in range(2, ws.max_row + 1)]
            items = [str(x).strip() for x in items if x is not None and str(x).strip()]
            if items:
                cols.append((column_key(items), items))
        tabs.append(cols)
    return tabs


def main():
    xlsx = Path(sys.argv[1])
    write = "--write" in sys.argv
    man = json.loads(BASE.read_text(encoding="utf-8"))
    old = OrderedDict((t["id"], t) for t in man["topics"])
    old_norm = {}
    for t in old.values():
        for k in (norm(t["name"]), norm(t.get("source", t["name"])), t["id"].replace("-", " ")):
            old_norm.setdefault(k, t["id"])

    tabs = read_tabs(xlsx)
    masters = {}
    for cols in tabs:
        for key, items in cols:
            if key not in masters or len(items) > len(masters[key]):
                masters[key] = items
    # every column must be a prefix of its master list
    for cols in tabs:
        for key, items in cols:
            assert masters[key][:len(items)] == items, f"column of {key} is not a prefix of its master list"

    # science rules: area -> rules
    sci_rules = {}
    for area, text in C.SCIENCE_AREAS.items():
        for k, (raw, dec) in parse_rules(text).items():
            if area == "=":
                sci_rules[k] = (raw, dec)
            elif area == "-":
                sci_rules[k] = (raw, "-")
            else:
                sci_rules[k] = (raw, f"{area}: {dec or display(raw)}")
    rules = {key: parse_rules(C.RULES.get(key, "")) for key in C.LIST_AREA}
    rules["sci"] = sci_rules

    new = OrderedDict()            # id -> topic
    where = {}                     # (key, norm raw) -> topic id or None
    first_pos = defaultdict(lambda: 10 ** 6)   # id -> best (lowest) position in any master list
    used_rules = defaultdict(set)
    dropped = []
    problems = []

    def add_new(name, area, raw):
        tid = slug(name)
        if tid in old:
            problems.append(f"new name {name!r} collides with existing id {tid} — mapped to the existing topic")
            return tid
        if tid in new:
            if new[tid]["category"] != area:
                problems.append(f"{name!r} assigned to both {new[tid]['category']} and {area}")
            if aka_ok(new[tid]["name"], display(raw)) and display(raw) not in new[tid]["aka"]:
                new[tid]["aka"].append(display(raw))
            return tid
        new[tid] = {"id": tid, "name": name, "category": area, "qbCategory": AREAS[area]["qb"],
                    "source": raw, "aka": [display(raw)] if aka_ok(name, display(raw)) else []}
        return tid

    for key in C.LIST_AREA:
        items = masters.get(key, [])
        default_area = C.LIST_AREA[key]
        math_from = items.index(C.BIOMATH_MATH_FROM) if key == "biomath" and C.BIOMATH_MATH_FROM in items else None
        for pos, raw in enumerate(items, 1):
            k = norm(raw)
            if (key, k) in where:
                tid = where[(key, k)]
            else:
                rule = rules[key].get(k)
                dec = None
                if rule:
                    used_rules[key].add(k)
                    dec = rule[1]
                if dec is None and k in old_norm and not rule:
                    tid = old_norm[k]
                elif dec == "-":
                    tid = None
                    dropped.append((LIST_LABEL[key], raw))
                elif dec and dec.startswith("="):
                    tid = dec[1:]
                    assert tid in old, f"unknown existing id {tid}"
                else:
                    area = default_area
                    if math_from is not None and pos - 1 >= math_from:
                        area = "Math"
                    name = dec or display(raw)
                    m = re.match(r"^([A-Z][A-Za-z ]+): (.+)$", name)
                    if m and m.group(1) in AREAS:
                        area, name = m.group(1), m.group(2)
                    if area is None:
                        problems.append(f"science entry without an area: {raw!r}")
                        tid = None
                    else:
                        tid = add_new(name, area, raw)
                where[(key, k)] = tid
            if tid and tid in new:
                first_pos[tid] = min(first_pos[tid], pos)

    for key in C.LIST_AREA:
        unused = set(rules[key]) - used_rules[key]
        for k in sorted(unused):
            problems.append(f"rule never matched in {key}: {rules[key][k][0]!r}")

    # order: old topics first within each area (unchanged relative order), then new by priority
    ordered = []
    for area in AREAS:
        ordered += [dict(t) for t in old.values() if t["category"] == area]
        ordered += sorted([t for t in new.values() if t["category"] == area], key=lambda t: first_pos[t["id"]])
    for i, t in enumerate(ordered, 1):
        t["order"] = i
        if "aka" in t and not t["aka"]:
            del t["aka"]
    counts = defaultdict(int)
    for t in ordered:
        counts[t["category"]] += 1

    # team study lists (anonymous): one per distinct tab assignment
    lists, seen = [], set()
    for cols in tabs:
        if not cols:
            continue
        parts, ids = [], []
        sig = []
        for key, items in cols:
            col_ids = []
            for raw in items:
                tid = where[(key, norm(raw))]
                if tid and tid not in col_ids:
                    col_ids.append(tid)
            full = []
            for raw in masters[key]:
                tid = where[(key, norm(raw))]
                if tid and tid not in full:
                    full.append(tid)
            label = LIST_LABEL[key]
            if key == "biomath":
                mf = masters[key].index(C.BIOMATH_MATH_FROM)
                if len(items) <= mf:
                    label = "Biology"
            if key == "biomath" and label == "Biology":
                bio_full = []
                for raw in masters[key][:masters[key].index(C.BIOMATH_MATH_FROM)]:
                    tid = where[(key, norm(raw))]
                    if tid and tid not in bio_full:
                        bio_full.append(tid)
                full = bio_full
            parts.append(f"{label} ({len(col_ids)})" if len(col_ids) == len(full) else f"{label} (first {len(col_ids)})")
            sig.append((key, len(items)))
            ids += [i for i in col_ids if i not in ids]
        sig = tuple(sig)
        if sig in seen:
            continue
        seen.add(sig)
        lists.append({"n": len(lists) + 1, "label": " + ".join(parts), "ids": ids})

    out = {
        "areas": [{"name": a, **meta, "count": counts[a]} for a, meta in AREAS.items()],
        "topics": ordered,
        "lists": lists,
    }

    # build queue: round-robin across areas by list position
    per_area = OrderedDict((a, sorted([t for t in new.values() if t["category"] == a], key=lambda t: first_pos[t["id"]])) for a in AREAS)
    queue = []
    while any(per_area.values()):
        for a in AREAS:
            if per_area[a]:
                t = per_area[a].pop(0)
                queue.append({"id": t["id"], "name": t["name"], "category": a, "pos": first_pos[t["id"]]})

    print(f"existing topics kept: {len(old)}   new topics: {len(new)}   total: {len(ordered)}")
    for a in AREAS:
        n_old = sum(1 for t in old.values() if t["category"] == a)
        print(f"  {a:18s} {counts[a]:4d}  ({n_old} existing + {counts[a] - n_old} new)")
    print(f"dropped rows: {len(dropped)}")
    print(f"study lists: {len(lists)}")
    for l in lists:
        print(f"  List {l['n']}: {l['label']} -> {len(l['ids'])} topics")
    if problems:
        print("\nPROBLEMS:")
        for p in problems:
            print("  " + p)

    if write:
        MANIFEST.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
        QUEUE.write_text(json.dumps(queue, ensure_ascii=False, indent=1), encoding="utf-8")
        lines = ["# Curation report: expanded study list (October 2026)", "",
                 f"{len(new)} new topics added; {len(old)} original topics unchanged. "
                 f"Rows dropped (junk, generic, duplicates of a dropped row, or unsuitable): {len(dropped)}.", "",
                 "## Dropped rows", ""]
        for lab, raw in dropped:
            why = C.LEFT_OUT_REASONS.get(raw, "not a study topic (junk row, fragment, or generic word)")
            lines.append(f"- {lab}: {raw} — {why}")
        lines += ["", "## Merged spellings (new topics with other names from the list)", ""]
        for t in new.values():
            if t.get("aka"):
                lines.append(f"- {t['name']} ← " + "; ".join(t["aka"]))
        REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
        print(f"\nwrote {MANIFEST.relative_to(ROOT)}, {QUEUE.relative_to(ROOT)}, {REPORT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
