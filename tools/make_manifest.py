"""Build data/manifest.json from the team's Study_List.xlsx.

The spreadsheet has one tab per student. Each column is a content area (header in row 1)
with that student's assigned topics below it. The website must not show student names, so
the manifest keeps only the topics and anonymous, numbered study lists (one per distinct
assignment).
"""
import json, re, sys, unicodedata
from collections import OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "research" / "Study_List.xlsx"
OUT = ROOT / "data" / "manifest.json"

# Spreadsheet header -> content area shown on the site
AREA_OF_HEADER = {
    "HISTORY": "American History",
    "HISTORY (European)": "European History",
    "RMPSS (Religion)": "Religion",
    "RELIGION": "Religion",
    "LIT": "Literature",
    "MYTHOLOGY": "Mythology",
    "SCIENCE": "Physics",
    "FINE ARTS": "Fine Arts",
    "FA": "Fine Arts",
    "SOCIAL SCIENCE": "Social Science",
    "PHILOSOPHY": "Philosophy",
    "EARTH SCIENCE": "Earth Science",
    "BIOLOGY": "Biology",
}

# Display order on the site, with the QBReader category each area maps to.
AREAS = OrderedDict([
    ("American History", {"qb": "History · American History", "short": "US History"}),
    ("European History", {"qb": "History · European History", "short": "Euro History"}),
    ("Literature", {"qb": "Literature", "short": "Literature"}),
    ("Mythology", {"qb": "Mythology", "short": "Mythology"}),
    ("Religion", {"qb": "Religion", "short": "Religion"}),
    ("Philosophy", {"qb": "Philosophy", "short": "Philosophy"}),
    ("Social Science", {"qb": "Social Science", "short": "Social Science"}),
    ("Fine Arts", {"qb": "Fine Arts", "short": "Fine Arts"}),
    ("Physics", {"qb": "Science · Physics", "short": "Physics"}),
    ("Biology", {"qb": "Science · Biology", "short": "Biology"}),
    ("Earth Science", {"qb": "Science · Other Science (Earth Science)", "short": "Earth Science"}),
])

# Cleaned display names for spreadsheet entries that need it.
DISPLAY = {
    "Magna Carta Libertatum": "Magna Carta",
    "Saint Joan of Arc": "Joan of Arc",
    "Baha'i Faith": "Bahá’í Faith",
    "Santeria": "Santería",
    "Oedipus": "Oedipus Rex",
    "The Catcher in The Rye": "The Catcher in the Rye",
    "Les Miserables": "Les Misérables",
    "A Midsummer Night's Dream": "A Midsummer Night’s Dream",
    "Women's Suffrage": "Women’s Suffrage",
    "Hundred Years' War": "Hundred Years’ War",
    "Thirty Years' War": "Thirty Years’ War",
    "Cu Chulainn": "Cú Chulainn",
    "Rene Descartes": "René Descartes",
    "Saint Paul's Cathedral": "St. Paul’s Cathedral",
    "DNA": "DNA",
}

# Ids already in use (keep them so saved progress carries over).
FIXED_IDS = {
    "Franklin Delano Roosevelt": "franklin-roosevelt",
}


def slug(s):
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[’'`.]", "", s)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s


def display_name(raw):
    raw = raw.strip()
    if raw in DISPLAY:
        return DISPLAY[raw]
    if raw and raw[0].islower():
        return raw[0].upper() + raw[1:]
    return raw


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    topics = OrderedDict()          # id -> topic
    area_order = {a: [] for a in AREAS}
    raw_lists = []                  # one per tab: [(area, [ids])]
    for ws in wb.worksheets:
        cols = []
        for c in range(1, ws.max_column + 1):
            header = ws.cell(1, c).value
            items = [ws.cell(r, c).value for r in range(2, ws.max_row + 1)]
            items = [str(x).strip() for x in items if x is not None and str(x).strip()]
            if not items:
                continue
            header = str(header).strip() if header else ""
            area = AREA_OF_HEADER.get(header)
            if not area:
                sys.exit(f"Unknown header {header!r} on a tab")
            ids = []
            for raw in items:
                name = display_name(raw)
                tid = FIXED_IDS.get(raw) or slug(name)
                if tid not in topics:
                    topics[tid] = {"id": tid, "name": name, "category": area, "qbCategory": AREAS[area]["qb"], "source": raw}
                    area_order[area].append(tid)
                elif topics[tid]["category"] != area:
                    sys.exit(f"{tid} appears under two areas")
                if tid not in ids:
                    ids.append(tid)
            cols.append((area, ids))
        if cols:
            raw_lists.append(cols)

    # Distinct assignments become numbered study lists (no names).
    lists, seen = [], set()
    for cols in raw_lists:
        key = tuple((a, tuple(ids)) for a, ids in cols)
        if key in seen:
            continue
        seen.add(key)
        parts = []
        for area, ids in cols:
            full = len(area_order[area])
            parts.append(f"{area} ({len(ids)})" if len(ids) == full else f"{area} (first {len(ids)})")
        lists.append({"n": len(lists) + 1, "label": " + ".join(parts),
                      "ids": [i for _, ids in cols for i in ids]})

    ordered = [topics[i] for a in AREAS for i in area_order[a]]
    for i, t in enumerate(ordered, 1):
        t["order"] = i
    data = {
        "areas": [{"name": a, **meta, "count": len(area_order[a])} for a, meta in AREAS.items()],
        "topics": ordered,
        "lists": lists,
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(ordered)} topics in {len(AREAS)} areas; {len(lists)} study lists")
    for a in AREAS:
        print(f"  {a}: {len(area_order[a])}")
    for l in lists:
        print(f"  List {l['n']}: {l['label']} -> {len(l['ids'])} topics")


if __name__ == "__main__":
    main()
